"""Exportação da lista de contatos para CSV e Excel."""

from __future__ import annotations

import csv
import re
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import COLUNAS, Lead

LARGURAS = {
    "nome": 34, "score": 8, "motivo_score": 46, "categoria": 22, "telefone": 18,
    "whatsapp": 16, "emails": 32, "site": 34, "instagram": 20, "facebook": 20,
    "linkedin": 20, "endereco": 46, "nota": 7, "avaliacoes": 12,
    "sinal_brasileiro": 9, "motivo_brasileiro": 42,
    "mensagem_abordagem": 60, "descricao_site": 46, "maps_url": 30,
    "query_origem": 28,
}

# Caracteres de controle que vêm colados no texto raspado dos sites (ESC de
# sequências ANSI, \x0b, \x0c, NUL...). O openpyxl recusa qualquer um deles com
# IllegalCharacterError, e um único lead sujo derrubava a exportação inteira.
# Quebra de linha e tabulação ficam: são válidas na célula e no CSV entre aspas.
RE_CONTROLE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f-\x9f]")


def _limpar(valor: Any) -> Any:
    if isinstance(valor, str):
        return RE_CONTROLE.sub("", valor)
    return valor


def _ordenar(leads: list[Lead]) -> list[Lead]:
    """Melhores leads primeiro; sem score, quem tem mais canal de contato.

    Quem foi marcado como público brasileiro sobe. Quando ninguém está marcado
    — o caso de qualquer busca dentro do Brasil — o campo vale 0 em todo mundo e
    a ordenação fica exatamente como era.
    """
    return sorted(
        leads,
        key=lambda l: (
            l.sinal_brasileiro >= 30,
            l.score if l.score is not None else -1,
            bool(l.emails) + bool(l.whatsapp) + bool(l.telefone),
            l.avaliacoes or 0,
        ),
        reverse=True,
    )


def _linhas(leads: list[Lead]) -> list[dict[str, Any]]:
    return [{k: _limpar(v) for k, v in l.to_row().items()} for l in leads]


def _preencher_aba(ws, linhas: list[dict], campos: list[str], titulos: list[str]) -> None:
    ws.append(titulos)

    cabecalho_fundo = PatternFill("solid", fgColor="1F3864")
    for celula in ws[1]:
        celula.font = Font(bold=True, color="FFFFFF")
        celula.fill = cabecalho_fundo
        celula.alignment = Alignment(vertical="center")

    for linha in linhas:
        ws.append([linha.get(c, "") for c in campos])

    for idx, campo in enumerate(campos, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = LARGURAS.get(campo, 20)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def exportar_csv(leads: list[Lead], destino: str | Path) -> Path:
    """Só o CSV (delimitador `;`, utf-8-sig para o Excel pt-BR abrir com acento)."""
    leads = _ordenar(leads)
    linhas = _linhas(leads)
    campos = [c for c, _ in COLUNAS]
    titulos = [t for _, t in COLUNAS]

    caminho = Path(destino).with_suffix(".csv")
    caminho.parent.mkdir(parents=True, exist_ok=True)
    with caminho.open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(titulos)
        for linha in linhas:
            w.writerow([linha.get(c, "") for c in campos])
    return caminho


def exportar_xlsx(leads: list[Lead], destino: str | Path) -> Path:
    """Só o Excel, com as abas `Leads`, `Com site` e `Sem site`."""
    leads = _ordenar(leads)
    campos = [c for c, _ in COLUNAS]
    titulos = [t for _, t in COLUNAS]

    caminho = Path(destino).with_suffix(".xlsx")
    caminho.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "Leads"
    _preencher_aba(ws, _linhas(leads), campos, titulos)

    ws_com_site = wb.create_sheet("Com site")
    _preencher_aba(ws_com_site, _linhas([l for l in leads if l.site]), campos, titulos)

    ws_sem_site = wb.create_sheet("Sem site")
    _preencher_aba(ws_sem_site, _linhas([l for l in leads if not l.site]), campos, titulos)

    wb.save(caminho)
    return caminho


def exportar(leads: list[Lead], destino: str) -> tuple[Path, Path]:
    """CSV e Excel juntos — o que a CLI usa. A web pede um de cada vez."""
    return exportar_csv(leads, destino), exportar_xlsx(leads, destino)
